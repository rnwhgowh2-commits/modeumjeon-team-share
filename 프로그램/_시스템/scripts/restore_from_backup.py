"""백업 (옛 정리된 옵션 832건) 복원 + 박스히어로 재고 sync.

이전 reset 작업이 옛 정리 결과 (색상 정리, 모델 분리 등) 을 날렸으므로 복원.
"""
import sys
import json
from datetime import datetime
from pathlib import Path
from sqlalchemy import text

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))


def main():
    backup_dir = sorted(Path(r'C:\dev\모음전 프로젝트\프로그램\_시스템\data').glob('backup_*'))[-1]
    print(f'백업: {backup_dir.name}')

    from shared.db import SessionLocal
    s = SessionLocal()
    try:
        # 1. 현재 옵션·의존 데이터 삭제 (잘못된 reset 결과)
        print('\n=== 현재 데이터 클린 ===')
        for tbl in ['inventory_txs', 'option_source_url_links', 'option_source_links',
                    'option_source_urls', 'option_price_config', 'etc_source_urls',
                    'option_account_registrations', 'option_product_links',
                    'option_benefit_overrides']:
            sp = s.begin_nested()
            try:
                r = s.execute(text(f'DELETE FROM {tbl}'))
                sp.commit()
                if r.rowcount > 0:
                    print(f'  {tbl}: {r.rowcount}건 삭제')
            except Exception as e:
                sp.rollback()
        r = s.execute(text('DELETE FROM options'))
        print(f'  options: {r.rowcount}건 삭제')
        s.commit()

        # 2. options 복원
        print('\n=== options.json 복원 ===')
        opts = json.loads((backup_dir / 'options.json').read_text(encoding='utf-8'))
        print(f'  복원 대상: {len(opts)}건')
        for o in opts:
            # 모든 컬럼 INSERT
            cols = ', '.join(o.keys())
            placeholders = ', '.join(f':{k}' for k in o.keys())
            sp = s.begin_nested()
            try:
                s.execute(text(f'INSERT INTO options ({cols}) VALUES ({placeholders})'), o)
                sp.commit()
            except Exception as e:
                sp.rollback()
                print(f'  실패: {o.get("canonical_sku")} - {e.__class__.__name__}')
        s.commit()

        new_total = s.execute(text('SELECT COUNT(*) FROM options')).scalar()
        print(f'  복원됨: {new_total}건')

        # 3. option_source_url_links 복원
        print('\n=== 매핑 복원 ===')
        mappings = json.loads((backup_dir / 'option_source_url_links.json').read_text(encoding='utf-8'))
        restored = 0
        for m in mappings:
            sp = s.begin_nested()
            try:
                s.execute(text(
                    'INSERT INTO option_source_url_links (option_canonical_sku, bundle_source_url_id, created_at) '
                    'VALUES (:s, :u, NOW())'
                ), {'s': m['option_canonical_sku'], 'u': m['bundle_source_url_id']})
                sp.commit()
                restored += 1
            except Exception:
                sp.rollback()
        s.commit()
        print(f'  복원됨: {restored}/{len(mappings)}건')

        # 4. option_product_links 복원
        print('\n=== option_product_links 복원 ===')
        plinks = json.loads((backup_dir / 'option_product_links.json').read_text(encoding='utf-8'))
        prestored = 0
        for p in plinks:
            sp = s.begin_nested()
            try:
                cols = ', '.join(p.keys())
                placeholders = ', '.join(f':{k}' for k in p.keys())
                s.execute(text(f'INSERT INTO option_product_links ({cols}) VALUES ({placeholders})'), p)
                sp.commit()
                prestored += 1
            except Exception:
                sp.rollback()
        s.commit()
        print(f'  복원됨: {prestored}/{len(plinks)}건')

        # 5. 박스히어로 재고 sync
        print('\n=== 박스히어로 재고 sync ===')
        import openpyxl
        bx_wb = openpyxl.load_workbook(r'C:\Users\seung\Downloads\Items_Export_99LAB_2026-05-27T22-48-42.xlsx', data_only=True)
        bx_ws = bx_wb['BoxHero']
        bx_hdr = [c.value for c in bx_ws[1]]
        bx_rows = [dict(zip(bx_hdr, r)) for r in bx_ws.iter_rows(min_row=2, values_only=True)]

        # 옵션 매핑 (boxhero_sku → canonical_sku)
        opt_by_bs = {}
        for r in s.execute(text(
            "SELECT boxhero_sku, canonical_sku FROM options WHERE boxhero_sku IS NOT NULL AND boxhero_sku != ''"
        )).fetchall():
            opt_by_bs[r[0]] = r[1]
        print(f'  옵션 boxhero_sku 매핑: {len(opt_by_bs)}')

        locs = s.execute(text(
            "SELECT id, name FROM inventory_locations WHERE deleted_at IS NULL"
        )).fetchall()
        loc_id_by_name = {r[1]: r[0] for r in locs}
        main_loc = loc_id_by_name.get('기본 위치')
        gros_loc = loc_id_by_name.get('그로스')

        ts = datetime.now()
        tx_inserted = 0
        for bx in bx_rows:
            sku = bx.get('SKU')
            if not sku:
                continue
            csku = opt_by_bs.get(sku)
            if not csku:
                continue
            qty_main = int(bx.get('수량(기본 위치)') or 0)
            qty_gros = int(bx.get('수량(그로스)') or 0)
            price = int(bx.get('구매가') or 0)
            if qty_main > 0:
                s.execute(text("""
                    INSERT INTO inventory_txs (option_canonical_sku, location_id, qty, tx_type, source, unit_purchase_price_at_tx, created_at, status)
                    VALUES (:cs, :loc, :qty, 'in', 'import', :price, :ts, 'completed')
                """), {'cs': csku, 'loc': main_loc, 'qty': qty_main, 'price': price, 'ts': ts})
                tx_inserted += 1
            if qty_gros > 0 and gros_loc:
                s.execute(text("""
                    INSERT INTO inventory_txs (option_canonical_sku, location_id, qty, tx_type, source, unit_purchase_price_at_tx, created_at, status)
                    VALUES (:cs, :loc, :qty, 'in', 'import', :price, :ts, 'completed')
                """), {'cs': csku, 'loc': gros_loc, 'qty': qty_gros, 'price': price, 'ts': ts})
                tx_inserted += 1
        s.commit()
        print(f'  inventory_txs: {tx_inserted}건')

        # 검증
        print('\n=== 검증 ===')
        from shared.inventory_stock import get_stock_batch
        all_skus = [r[0] for r in s.execute(text('SELECT canonical_sku FROM options')).fetchall()]
        stock = get_stock_batch(s, all_skus)
        stock_total = sum(stock.values())
        total = s.execute(text('SELECT COUNT(*) FROM options')).scalar()
        active = s.execute(text('SELECT COUNT(*) FROM options WHERE is_active = true')).scalar()
        print(f'옵션: {total} (활성 {active})')
        print(f'재고: {stock_total} (목표 722)')
    finally:
        s.close()


if __name__ == '__main__':
    main()
